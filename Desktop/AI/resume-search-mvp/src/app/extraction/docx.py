from io import BytesIO

from app.extraction.base import ExtractionResult


class DocxTextExtractor:
    def extract(self, file_bytes: bytes) -> ExtractionResult:
        try:
            from docx import Document

            document = Document(BytesIO(file_bytes))
            text_parts: list[str] = []

            for paragraph in document.paragraphs:
                text_parts.append(paragraph.text)

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            text_parts.append(paragraph.text)

            cleaned_text = "\n".join(
                text.strip() for text in text_parts if text.strip()
            )

            if not cleaned_text:
                return ExtractionResult(text="", status="empty_text")

            return ExtractionResult(text=cleaned_text, status="extracted")
        except Exception:
            return ExtractionResult(text="", status="failed")
