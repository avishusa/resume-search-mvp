from io import BytesIO

from docx import Document

from app.extraction.docx import DocxTextExtractor


def _build_docx_bytes(
    paragraphs: list[str] | None = None,
    table_values: list[list[str]] | None = None,
) -> bytes:
    document = Document()

    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)

    if table_values:
        table = document.add_table(rows=len(table_values), cols=len(table_values[0]))
        for row_index, row_values in enumerate(table_values):
            for column_index, value in enumerate(row_values):
                table.cell(row_index, column_index).text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_extractor_extracts_paragraph_text() -> None:
    extractor = DocxTextExtractor()

    result = extractor.extract(
        _build_docx_bytes(paragraphs=["AI Engineer", "Python FastAPI"])
    )

    assert result.status == "extracted"
    assert result.text == "AI Engineer\nPython FastAPI"


def test_docx_extractor_extracts_table_cell_text() -> None:
    extractor = DocxTextExtractor()

    result = extractor.extract(
        _build_docx_bytes(table_values=[["Skill", "Python"], ["Tool", "Docker"]])
    )

    assert result.status == "extracted"
    assert "Skill" in result.text
    assert "Python" in result.text
    assert "Docker" in result.text


def test_docx_extractor_returns_empty_text_for_blank_docx() -> None:
    extractor = DocxTextExtractor()

    result = extractor.extract(_build_docx_bytes())

    assert result.status == "empty_text"
    assert result.text == ""


def test_docx_extractor_returns_failed_for_invalid_docx_bytes() -> None:
    extractor = DocxTextExtractor()

    result = extractor.extract(b"not a docx")

    assert result.status == "failed"
    assert result.text == ""


def test_docx_extractor_does_not_return_not_implemented() -> None:
    extractor = DocxTextExtractor()

    result = extractor.extract(_build_docx_bytes(paragraphs=["AI Engineer"]))

    assert result.status != "not_implemented"
